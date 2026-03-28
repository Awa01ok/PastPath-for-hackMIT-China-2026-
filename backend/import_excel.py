import sqlite3
from pathlib import Path
from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = BASE_DIR.parent
DEFAULT_XLSX = PROJECT_ROOT / 'database' / 'pastpath_database_template.xlsx'
DEFAULT_MAP_XLSX = PROJECT_ROOT / 'database' / 'pastpath_map_database.xlsx'
DEFAULT_DB = BASE_DIR / 'pastpath.db'


def parse_bool(value):
    if value is None:
        return 0
    s = str(value).strip().lower()
    return 1 if s in {'1', 'true', 'yes', 'y', 'enabled'} else 0


def parse_text(value):
    if value is None:
        return None
    s = str(value).strip()
    return s if s else None


def parse_float(value, default=None):
    if value is None or value == '':
        return default
    try:
        return float(value)
    except Exception:
        return default


def parse_int(value, default=None):
    if value is None or value == '':
        return default
    try:
        return int(value)
    except Exception:
        return default


def split_multi(value):
    if value is None:
        return []
    s = str(value).strip()
    if not s:
        return []
    normalized = s.replace('；', ';').replace('，', ',').replace('|', ',')
    parts = []
    for chunk in normalized.split(';'):
        for item in chunk.split(','):
            item = item.strip()
            if item:
                parts.append(item)
    seen = set()
    result = []
    for item in parts:
        key = item.lower()
        if key not in seen:
            seen.add(key)
            result.append(item)
    return result


def create_schema(conn):
    conn.executescript(
        """
        PRAGMA foreign_keys = ON;

        DROP TABLE IF EXISTS event_sources;
        DROP TABLE IF EXISTS event_tags;
        DROP TABLE IF EXISTS person_tags;
        DROP TABLE IF EXISTS events;
        DROP TABLE IF EXISTS people;
        DROP TABLE IF EXISTS timeline_layer_map;
        DROP TABLE IF EXISTS layer_registry;
        DROP TABLE IF EXISTS timelines;
        DROP TABLE IF EXISTS territories;

        CREATE TABLE people (
            person_id TEXT PRIMARY KEY,
            name TEXT NOT NULL,
            display_years TEXT,
            summary TEXT,
            portrait_local TEXT,
            qid TEXT,
            wikipedia_url TEXT,
            sort_order INTEGER DEFAULT 9999,
            enabled INTEGER DEFAULT 1
        );

        CREATE TABLE person_tags (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            person_id TEXT NOT NULL,
            tag_group TEXT NOT NULL,
            tag_value TEXT NOT NULL,
            FOREIGN KEY (person_id) REFERENCES people(person_id) ON DELETE CASCADE
        );

        CREATE TABLE events (
            event_id TEXT PRIMARY KEY,
            person_id TEXT NOT NULL,
            title TEXT NOT NULL,
            description TEXT,
            date_start TEXT,
            date_end TEXT,
            year INTEGER,
            place TEXT,
            lng REAL,
            lat REAL,
            zoom REAL DEFAULT 5,
            pitch REAL DEFAULT 20,
            bearing REAL DEFAULT 0,
            image_local TEXT,
            route_order INTEGER DEFAULT 9999,
            timeline_key TEXT,
            territory_key TEXT,
            enabled INTEGER DEFAULT 1,
            FOREIGN KEY (person_id) REFERENCES people(person_id) ON DELETE CASCADE
        );

        CREATE TABLE event_sources (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            event_id TEXT NOT NULL,
            source_url TEXT NOT NULL,
            FOREIGN KEY (event_id) REFERENCES events(event_id) ON DELETE CASCADE
        );

        CREATE TABLE event_tags (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            event_id TEXT NOT NULL,
            tag_value TEXT NOT NULL,
            FOREIGN KEY (event_id) REFERENCES events(event_id) ON DELETE CASCADE
        );

        CREATE TABLE timelines (
            timeline_key TEXT PRIMARY KEY,
            time_key TEXT,
            year INTEGER,
            display_label TEXT,
            label TEXT,
            start_date TEXT,
            end_date TEXT,
            region_scope TEXT,
            territory_key TEXT,
            sort_order INTEGER DEFAULT 9999,
            is_global INTEGER DEFAULT 1,
            enabled INTEGER DEFAULT 1
        );

        CREATE TABLE territories (
            territory_key TEXT PRIMARY KEY,
            layer_key TEXT,
            layer_type TEXT,
            region_scope TEXT,
            time_key TEXT,
            year INTEGER,
            label TEXT,
            geojson_file TEXT,
            style_fill TEXT,
            style_opacity REAL,
            line_color TEXT,
            line_width REAL,
            source_name TEXT,
            source_license TEXT,
            source_url TEXT,
            priority INTEGER DEFAULT 9999,
            enabled INTEGER DEFAULT 1,
            notes TEXT
        );

        CREATE TABLE layer_registry (
            layer_key TEXT PRIMARY KEY,
            layer_name TEXT,
            layer_type TEXT,
            default_visible INTEGER DEFAULT 1,
            z_index INTEGER DEFAULT 10,
            description TEXT,
            enabled INTEGER DEFAULT 1
        );

        CREATE TABLE timeline_layer_map (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            timeline_key TEXT NOT NULL,
            region_scope TEXT,
            country_territory_key TEXT,
            control_territory_key TEXT,
            camp_territory_key TEXT,
            frontline_territory_key TEXT,
            fallback_mode TEXT,
            notes TEXT
        );

        CREATE INDEX idx_people_enabled ON people(enabled, sort_order, name);
        CREATE INDEX idx_events_person ON events(person_id, enabled, route_order, date_start);
        CREATE INDEX idx_person_tags_person ON person_tags(person_id, tag_group);
        CREATE INDEX idx_event_tags_event ON event_tags(event_id);
        CREATE INDEX idx_event_sources_event ON event_sources(event_id);
        CREATE INDEX idx_timelines_year ON timelines(year, sort_order);
        CREATE INDEX idx_territories_year ON territories(year);
        CREATE INDEX idx_territories_region_layer ON territories(region_scope, layer_type, year);
        CREATE INDEX idx_timeline_layer_map_timeline ON timeline_layer_map(timeline_key, region_scope);
        """
    )
    conn.commit()


def import_people(ws, conn):
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(h).strip() if h is not None else '' for h in rows[0]]
    for raw_row in rows[1:]:
        row = dict(zip(headers, raw_row))
        person_id = parse_text(row.get('person_id'))
        if not person_id:
            continue

        enabled = parse_bool(row.get('enabled'))
        conn.execute(
            """
            INSERT INTO people (
                person_id, name, display_years, summary, portrait_local,
                qid, wikipedia_url, sort_order, enabled
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                person_id,
                parse_text(row.get('name')) or person_id,
                parse_text(row.get('display_years')),
                parse_text(row.get('summary')),
                parse_text(row.get('portrait_local')),
                parse_text(row.get('qid')),
                parse_text(row.get('wikipedia_url')),
                parse_int(row.get('sort_order'), 9999),
                enabled,
            ),
        )

        tag_map = {
            'roles': split_multi(row.get('roles')),
            'domains': split_multi(row.get('domains')),
            'countries': split_multi(row.get('countries')),
            'theaters': split_multi(row.get('theaters')),
            'themes': split_multi(row.get('themes')),
            'eras': split_multi(row.get('eras')),
        }

        for tag_group, values in tag_map.items():
            for value in values:
                conn.execute(
                    "INSERT INTO person_tags (person_id, tag_group, tag_value) VALUES (?, ?, ?)",
                    (person_id, tag_group, value),
                )


def import_events(ws, conn):
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(h).strip() if h is not None else '' for h in rows[0]]

    for raw_row in rows[1:]:
        row = dict(zip(headers, raw_row))
        event_id_raw = parse_text(row.get('event_id'))
        person_id = parse_text(row.get('person_id'))
        if not event_id_raw or not person_id:
            continue

        event_id = f"{person_id}__{event_id_raw}"
        enabled = parse_bool(row.get('enabled'))

        conn.execute(
            """
            INSERT INTO events (
                event_id, person_id, title, description, date_start, date_end, year,
                place, lng, lat, zoom, pitch, bearing, image_local,
                route_order, timeline_key, territory_key, enabled
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                event_id,
                person_id,
                parse_text(row.get('title')) or event_id_raw,
                parse_text(row.get('description')),
                parse_text(row.get('date_start')),
                parse_text(row.get('date_end')),
                parse_int(row.get('year')),
                parse_text(row.get('place')),
                parse_float(row.get('lng')),
                parse_float(row.get('lat')),
                parse_float(row.get('zoom'), 5),
                parse_float(row.get('pitch'), 20),
                parse_float(row.get('bearing'), 0),
                parse_text(row.get('image_local')),
                parse_int(row.get('route_order'), 9999),
                parse_text(row.get('timeline_key')),
                parse_text(row.get('territory_key')),
                enabled,
            ),
        )

        for url in split_multi(row.get('source_urls')):
            conn.execute("INSERT INTO event_sources (event_id, source_url) VALUES (?, ?)", (event_id, url))
        for tag in split_multi(row.get('event_tags')):
            conn.execute("INSERT INTO event_tags (event_id, tag_value) VALUES (?, ?)", (event_id, tag))


def import_timeline(ws, conn):
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(h).strip() if h is not None else '' for h in rows[0]]
    for raw_row in rows[1:]:
        row = dict(zip(headers, raw_row))
        timeline_key = parse_text(row.get('timeline_key'))
        if not timeline_key:
            continue
        conn.execute(
            """
            INSERT INTO timelines (
                timeline_key, time_key, year, display_label, label, start_date, end_date,
                region_scope, territory_key, sort_order, is_global, enabled
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                timeline_key,
                parse_text(row.get('time_key')),
                parse_int(row.get('year')),
                parse_text(row.get('display_label')),
                parse_text(row.get('label')),
                parse_text(row.get('start_date')),
                parse_text(row.get('end_date')),
                parse_text(row.get('region_scope')),
                parse_text(row.get('territory_key')),
                parse_int(row.get('sort_order'), 9999),
                parse_bool(row.get('is_global')),
                parse_bool(row.get('enabled')),
            ),
        )


def import_territories(ws, conn):
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(h).strip() if h is not None else '' for h in rows[0]]
    for raw_row in rows[1:]:
        row = dict(zip(headers, raw_row))
        territory_key = parse_text(row.get('territory_key'))
        if not territory_key:
            continue
        conn.execute(
            """
            INSERT INTO territories (
                territory_key, layer_key, layer_type, region_scope, time_key, year, label,
                geojson_file, style_fill, style_opacity, line_color, line_width,
                source_name, source_license, source_url, priority, enabled, notes
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                territory_key,
                parse_text(row.get('layer_key')),
                parse_text(row.get('layer_type')),
                parse_text(row.get('region_scope')),
                parse_text(row.get('time_key')),
                parse_int(row.get('year')),
                parse_text(row.get('label')),
                parse_text(row.get('geojson_file')),
                parse_text(row.get('style_fill')),
                parse_float(row.get('style_opacity'), 0.35),
                parse_text(row.get('line_color')),
                parse_float(row.get('line_width'), 1),
                parse_text(row.get('source_name')),
                parse_text(row.get('source_license')),
                parse_text(row.get('source_url')),
                parse_int(row.get('priority'), 9999),
                parse_bool(row.get('enabled')),
                parse_text(row.get('notes')),
            ),
        )


def import_layer_registry(ws, conn):
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(h).strip() if h is not None else '' for h in rows[0]]
    for raw_row in rows[1:]:
        row = dict(zip(headers, raw_row))
        layer_key = parse_text(row.get('layer_key'))
        if not layer_key:
            continue
        conn.execute(
            """
            INSERT INTO layer_registry (
                layer_key, layer_name, layer_type, default_visible, z_index, description, enabled
            ) VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (
                layer_key,
                parse_text(row.get('layer_name')),
                parse_text(row.get('layer_type')),
                parse_bool(row.get('default_visible')),
                parse_int(row.get('z_index'), 10),
                parse_text(row.get('description')),
                parse_bool(row.get('enabled')),
            ),
        )


def import_timeline_layer_map(ws, conn):
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(h).strip() if h is not None else '' for h in rows[0]]
    for raw_row in rows[1:]:
        row = dict(zip(headers, raw_row))
        timeline_key = parse_text(row.get('timeline_key'))
        if not timeline_key:
            continue
        conn.execute(
            """
            INSERT INTO timeline_layer_map (
                timeline_key, region_scope, country_territory_key, control_territory_key,
                camp_territory_key, frontline_territory_key, fallback_mode, notes
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                timeline_key,
                parse_text(row.get('region_scope')),
                parse_text(row.get('country_territory_key')),
                parse_text(row.get('control_territory_key')),
                parse_text(row.get('camp_territory_key')),
                parse_text(row.get('frontline_territory_key')),
                parse_text(row.get('fallback_mode')),
                parse_text(row.get('notes')),
            ),
        )


def load_required_workbook(path):
    if not path.exists():
        raise FileNotFoundError(f'Excel file not found: {path}')
    return load_workbook(path, data_only=True)


def main(xlsx_path=DEFAULT_XLSX, map_xlsx_path=DEFAULT_MAP_XLSX, db_path=DEFAULT_DB):
    xlsx_path = Path(xlsx_path)
    map_xlsx_path = Path(map_xlsx_path)
    db_path = Path(db_path)

    story_wb = load_required_workbook(xlsx_path)
    required_story = {'people', 'events'}
    missing_story = required_story - set(story_wb.sheetnames)
    if missing_story:
        raise ValueError(f'Missing sheets in story workbook: {sorted(missing_story)}')

    map_wb = load_required_workbook(map_xlsx_path) if map_xlsx_path.exists() else story_wb
    required_map = {'timeline', 'territories'}
    missing_map = required_map - set(map_wb.sheetnames)
    if missing_map:
        raise ValueError(f'Missing sheets in map workbook: {sorted(missing_map)}')

    conn = sqlite3.connect(db_path)
    try:
        create_schema(conn)
        import_people(story_wb['people'], conn)
        import_events(story_wb['events'], conn)
        import_timeline(map_wb['timeline'], conn)
        import_territories(map_wb['territories'], conn)
        if 'layer_registry' in map_wb.sheetnames:
            import_layer_registry(map_wb['layer_registry'], conn)
        if 'timeline_layer_map' in map_wb.sheetnames:
            import_timeline_layer_map(map_wb['timeline_layer_map'], conn)
        conn.commit()
    finally:
        conn.close()

    print(f'Imported story workbook: {xlsx_path}')
    print(f'Imported map workbook: {map_xlsx_path if map_wb is not story_wb else xlsx_path}')
    print(f'Imported Excel into SQLite: {db_path}')


if __name__ == '__main__':
    main()
